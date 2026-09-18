from pathlib import Path
from docx import Document
from docx.shared import Inches

root = Path(r"C:\Users\cj_wo\OneDrive\Desktop\IoT-Digital-Repository")
out_file = root / "jerrydfirstdelivery.docx"
img_path = root / "diagrams" / "architecture.png"
sample_path = root / "samples" / "phase2_sample_output.txt"

report = Document()
report.styles["Normal"].font.name = "Calibri"

report.add_heading("GreenHouseWatch — Phase 1 Design and Phase 2 Simulation Delivery", 0)
report.add_paragraph("Course: CTEC 651 – Internet Technologies Discovery")
report.add_paragraph("Instructor: Prof. F. Njeh")
report.add_paragraph("Student / Team: IPMC BIT")
report.add_paragraph("Submission focus: Phase 1 Design (Due 7 Sep 2026) and Phase 2 Simulation (Due 14 Sep 2026)")

report.add_heading("1. Project overview", level=1)
report.add_paragraph(
    "GreenHouseWatch is a local IoT digital repository designed for greenhouse monitoring in a low-cost, free-to-use environment. "
    "The system captures temperature, humidity, soil moisture, and light readings from a virtual greenhouse sensor, publishes them through MQTT, "
    "records them in a local SQLite database, and supports analysis and visualisation without relying on paid cloud services."
)

report.add_heading("2. Problem statement and use case", level=1)
report.add_paragraph(
    "Small greenhouse and urban-farm operators lose plants when temperature, humidity, or soil moisture drift outside a safe operating band, often while no one is actively monitoring a thermometer or dashboard. "
    "Commercial IoT platforms are usually paid, cloud-locked, and too heavy for a class-scale system. GreenHouseWatch addresses this by focusing on local capture, storage, and inspection of greenhouse sensor data using free tools and simple automation."
)
report.add_paragraph(
    "The use case simulates a single greenhouse node, GH-SENSOR-01, that emits a JSON reading every two seconds. Each reading contains device metadata and environmental values, and the system stores the historical stream for later interpretation and alerting."
)

report.add_heading("3. Design specification", level=1)
report.add_paragraph(
    "The system follows a standard IoT flow: sensor simulator → MQTT broker → subscriber → SQLite storage → visualisation/export. "
    "The design uses the public HiveMQ broker and stores readings in a local repository to enable later charting and review."
)
if img_path.exists():
    report.add_paragraph("Architecture diagram:")
    report.add_picture(str(img_path), width=Inches(6.0))

report.add_paragraph(
    "Data fields used in the project include timestamp, ISO timestamp, device ID, temperature_c, humidity_pct, soil_moisture_pct, light_lux, alert_flag, and alert_reason. "
    "The alert logic flags readings that breach the safe bands: temperature 16–32 °C, humidity 40–80 %, and soil moisture ≥ 30 %."
)

report.add_heading("4. Implementation summary", level=1)
report.add_paragraph(
    "The design is implemented in the repository with the following components: docs/PHASE1_Project_Proposal.md, diagrams/architecture.png, src/simulator.py, and supporting source modules for MQTT, SQLite, and visualization. "
    "The project uses Python, paho-mqtt, SQLite, Matplotlib, and optional Streamlit for dashboarding."
)

report.add_heading("5. Simulation deliverable", level=1)
report.add_paragraph(
    "The Phase 2 simulation is implemented in src/simulator.py. It generates realistic greenhouse readings with day/night dynamics, light and temperature variation, humidity inversions, soil moisture decay, and injected anomaly conditions. "
    "This provides a credible time-series stream suitable for later database storage and alert interpretation."
)

if sample_path.exists():
    sample_lines = sample_path.read_text(encoding="utf-8").splitlines()
    for line in sample_lines[:8]:
        if line.strip():
            report.add_paragraph(line)

report.add_paragraph(
    "The sample output confirms the expected data format: timestamp, iso_time, device_id, temperature_c, humidity_pct, soil_moisture_pct, light_lux, alert_flag, and alert_reason. "
    "The simulator also demonstrates alert conditions such as high humidity and other threshold breaches in the generated readings."
)

report.add_heading("6. Submission checklist", level=1)
report.add_paragraph("Phase 1 Design — Due 7 Sep 2026 — Submit docs/PHASE1_Project_Proposal.md (or .docx) and diagrams/architecture.png")
report.add_paragraph("Phase 2 Simulation — Due 14 Sep 2026 — Submit src/simulator.py and samples/phase2_sample_output.txt")

report.add_heading("7. Conclusion", level=1)
report.add_paragraph(
    "GreenHouseWatch demonstrates a practical local IoT repository pattern for smart greenhouse monitoring. The project balances a clear architectural design with a realistic simulation that can later feed live MQTT transmission, SQLite storage, and dashboard analysis without requiring paid cloud services."
)

report.save(out_file)
print(f"Created: {out_file}")
print(f"File size: {out_file.stat().st_size} bytes")
